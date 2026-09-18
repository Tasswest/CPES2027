#!/usr/bin/env python3
"""Génère l'article au format Word, destiné à être retravaillé à la main.

Le texte vient de `article_contenu.py`, partagé avec la version PDF : les deux
documents sont identiques au moment de la génération. Une fois le Word modifié à
la main, c'est lui qui fait foi — relancer ce script écraserait ces modifications.

Pour que le document reste agréable à éditer, la mise en forme passe par des
styles nommés plutôt que par du formatage direct (changer la police du style
« Normal » se propage à tout le corps du texte), et les légendes sont numérotées
par des champs Word (SEQ) : insérer une figure renumérote les suivantes.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from article_contenu import construire

OUT = Path("article_ziak_stylometrie.docx")
IMAGES = Path("images")
POLICE = "Times New Roman"
ACCENT = RGBColor(0x2F, 0x6F, 0x9F)
GRIS = RGBColor(0x55, 0x55, 0x55)
LARGEUR_TEXTE_CM = 16.6

BALISE = re.compile(r"(<b>|</b>|<i>|</i>|<sub>|</sub>|<super>|</super>|<br/>)")


def entites(texte: str) -> str:
    return (texte.replace("&nbsp;", " ").replace("&lt;", "<")
            .replace("&gt;", ">").replace("&amp;", "&"))


def ecrire(paragraphe, texte: str) -> None:
    """Traduit le balisage minimal de l'article en runs Word."""
    gras = ital = indice = expo = 0
    for morceau in BALISE.split(texte):
        if not morceau:
            continue
        if morceau == "<br/>":
            paragraphe.add_run().add_break()
        elif morceau in ("<b>", "</b>"):
            gras += 1 if morceau == "<b>" else -1
        elif morceau in ("<i>", "</i>"):
            ital += 1 if morceau == "<i>" else -1
        elif morceau in ("<sub>", "</sub>"):
            indice += 1 if morceau == "<sub>" else -1
        elif morceau in ("<super>", "</super>"):
            expo += 1 if morceau == "<super>" else -1
        else:
            run = paragraphe.add_run(entites(morceau))
            if gras > 0:
                run.bold = True
            if ital > 0:
                run.italic = True
            if indice > 0:
                run.font.subscript = True
            if expo > 0:
                run.font.superscript = True


def xml(tag: str, **attributs) -> OxmlElement:
    el = OxmlElement(tag)
    for k, v in attributs.items():
        el.set(qn(k), v)
    return el


def police(style, taille, gras=None, italique=None, couleur=None) -> None:
    """Fixe la police d'un style en neutralisant les polices du thème Word."""
    style.font.name = POLICE
    style.font.size = Pt(taille)
    if gras is not None:
        style.font.bold = gras
    if italique is not None:
        style.font.italic = italique
    if couleur is not None:
        style.font.color.rgb = couleur
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(att), None)
    rfonts.set(qn("w:eastAsia"), POLICE)


def paragraphe_style(doc, nom, base="Normal"):
    styles = doc.styles
    if nom in [s.name for s in styles]:
        return styles[nom]
    st = styles.add_style(nom, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = styles[base]
    st.quick_style = True
    return st


def champ(paragraphe, instruction: str, valeur: str, gras: bool = False) -> None:
    """Insère un champ Word (SEQ, PAGE) avec sa valeur déjà calculée."""
    fld = xml("w:fldSimple", **{"w:instr": f" {instruction} "})
    r = OxmlElement("w:r")
    if gras:
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:b"))
        r.append(rpr)
    t = OxmlElement("w:t")
    t.text = valeur
    r.append(t)
    fld.append(r)
    paragraphe._p.append(fld)


def styles(doc) -> None:
    normal = doc.styles["Normal"]
    police(normal, 11)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nom, taille, couleur, avant, apres in [("Heading 1", 14, ACCENT, 16, 6),
                                                ("Heading 2", 12, RGBColor(0, 0, 0), 10, 4)]:
        st = doc.styles[nom]
        police(st, taille, gras=True, italique=False, couleur=couleur)
        st.paragraph_format.space_before = Pt(avant)
        st.paragraph_format.space_after = Pt(apres)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    reglages = {
        "Titre de l'article": dict(taille=20, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER, apres=4),
        "Sous-titre de l'article": dict(taille=12.5, italique=True, couleur=GRIS,
                                        align=WD_ALIGN_PARAGRAPH.CENTER, apres=14),
        "Auteur": dict(taille=11, align=WD_ALIGN_PARAGRAPH.CENTER, apres=2),
        "Affiliation": dict(taille=9.5, couleur=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, apres=16),
        "Résumé": dict(taille=10, retrait=0.8, apres=5),
        "Note": dict(taille=10, retrait=0.4, apres=8, avant=4),
        "En clair": dict(taille=10, retrait=0.4, apres=8, avant=4),
        "Référence": dict(taille=9.5, apres=3),
        "Cellule": dict(taille=9.5, apres=0, interligne=1.0),
    }
    for nom, r in reglages.items():
        st = paragraphe_style(doc, nom)
        police(st, r["taille"], gras=r.get("gras"), italique=r.get("italique"),
               couleur=r.get("couleur"))
        f = st.paragraph_format
        f.space_after = Pt(r.get("apres", 6))
        f.space_before = Pt(r.get("avant", 0))
        if "align" in r:
            f.alignment = r["align"]
        if "retrait" in r:
            f.left_indent = f.right_indent = Cm(r["retrait"])
        if "interligne" in r:
            f.line_spacing = r["interligne"]

    ref = doc.styles["Référence"].paragraph_format
    ref.left_indent, ref.first_line_indent = Cm(0.6), Cm(-0.6)

    # Le schéma impose l'ordre des enfants de w:pPr : bordure, puis ombrage,
    # tous deux avant l'espacement et les retraits.
    apres = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
             "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
             "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
             "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
             "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
             "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    # Les réserves méthodologiques sont bleu-gris, les encadrés « En clair »
    # ambrés : la couleur seule doit suffire à les distinguer en feuilletant.
    for nom, filet, fond in [("Note", "2F6F9F", "F4F6F8"),
                             ("En clair", "D9A441", "FDF6E7")]:
        ppr = doc.styles[nom].element.get_or_add_pPr()
        bordure = xml("w:pBdr")
        bordure.append(xml("w:left", **{"w:val": "single", "w:sz": "18",
                                        "w:space": "8", "w:color": filet}))
        ppr.insert_element_before(bordure, "w:shd", *apres)
        ppr.insert_element_before(
            xml("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fond}),
            *apres)

    legende = doc.styles["Caption"]
    police(legende, 9, gras=False, italique=False, couleur=GRIS)
    legende.paragraph_format.space_after = Pt(12)
    legende.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


class RenduWord:
    STYLES = {"titre": "Titre de l'article", "soustitre": "Sous-titre de l'article",
              "auteur": "Auteur", "date": "Affiliation", "abstract": "Résumé",
              "h1": "Heading 1", "h2": "Heading 2", "p": "Normal", "note": "Note",
              "clair": "En clair", "ref": "Référence"}

    def __init__(self, doc):
        self.doc = doc
        self.n_fig = self.n_tab = 0
        self.dernier = None

    def p(self, texte, style="p"):
        par = self.doc.add_paragraph(style=self.STYLES[style])
        ecrire(par, texte)
        self.dernier = par

    def gap(self, h=6):
        if self.dernier is not None:
            pf = self.dernier.paragraph_format
            pf.space_after = Pt((pf.space_after.pt if pf.space_after else 6) + h)

    def legende(self, prefixe, n, texte, garder_avec_suivant=False):
        par = self.doc.add_paragraph(style="Caption")
        par.add_run(f"{prefixe} ").bold = True
        champ(par, f"SEQ {prefixe} \\* ARABIC", str(n), gras=True)
        par.add_run(". ").bold = True
        ecrire(par, texte)
        par.paragraph_format.keep_with_next = garder_avec_suivant
        self.dernier = par

    def figure(self, image, legende):
        self.n_fig += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_before = Pt(6)
        par.add_run().add_picture(str(IMAGES / image), width=Cm(LARGEUR_TEXTE_CM))
        self.legende("Figure", self.n_fig, legende)

    def tableau(self, lignes, legende, widths=None, gauche=(0,)):
        self.n_tab += 1
        self.legende("Tableau", self.n_tab, legende, garder_avec_suivant=True)
        n_col = len(lignes[0])
        widths = widths or [LARGEUR_TEXTE_CM / n_col] * n_col
        t = self.doc.add_table(rows=len(lignes), cols=n_col)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        # Word lit la grille, pas seulement la largeur des cellules : sans cette
        # reprise, toutes les colonnes ressortent égales quelle que soit `widths`.
        tblw = t._tbl.tblPr.find(qn("w:tblW"))
        if tblw is not None:
            tblw.set(qn("w:type"), "dxa")
            tblw.set(qn("w:w"), str(int(Cm(sum(widths)).twips)))
        for col, largeur in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(Cm(largeur).twips)))
        tpr = t._tbl.tblPr
        bords = xml("w:tblBorders")
        for cote in ("top", "bottom", "insideH"):
            bords.append(xml(f"w:{cote}", **{"w:val": "single", "w:sz": "4",
                                             "w:space": "0", "w:color": "C8CDD2"}))
        # Ordre imposé dans w:tblPr : les bordures précèdent w:tblLayout, que
        # python-docx a déjà posé (mise en page fixe) via `autofit = False`.
        tpr.insert_element_before(bords, "w:shd", "w:tblLayout", "w:tblCellMar",
                                  "w:tblLook", "w:tblCaption", "w:tblDescription",
                                  "w:tblPrChange")

        for i, ligne in enumerate(lignes):
            row = t.rows[i]
            trpr = row._tr.get_or_add_trPr()
            trpr.append(xml("w:cantSplit"))
            if i == 0:
                trpr.append(xml("w:tblHeader"))
            for j, valeur in enumerate(ligne):
                cell = row.cells[j]
                cell.width = Cm(widths[j])
                tcpr = cell._tc.get_or_add_tcPr()
                fond = "2F6F9F" if i == 0 else ("F6F8FA" if i % 2 == 0 else None)
                if fond:
                    tcpr.append(xml("w:shd", **{"w:val": "clear", "w:color": "auto",
                                                "w:fill": fond}))
                tcpr.append(xml("w:vAlign", **{"w:val": "center"}))
                par = cell.paragraphs[0]
                par.style = self.doc.styles["Cellule"]
                par.alignment = (WD_ALIGN_PARAGRAPH.LEFT if (j in gauche and i > 0)
                                 else WD_ALIGN_PARAGRAPH.CENTER)
                if i == 0:
                    ecrire(par, f"<b>{valeur}</b>")
                    for run in par.runs:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    ecrire(par, str(valeur))
        # Espace après le tableau, porté par un paragraphe vide discret.
        vide = self.doc.add_paragraph()
        vide.paragraph_format.space_after = Pt(4)
        self.dernier = vide


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)
    styles(doc)

    pied = section.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    champ(pied, "PAGE", "1")

    props = doc.core_properties
    props.title = "Qui se cache derrière Ziak ? Une enquête stylométrique"
    props.author = "Tassilo Westphalen"
    props.subject = "Stylométrie et attribution d'auteur sur le corpus LRFAF"
    props.keywords = "stylométrie, attribution d'auteur, rap français, LRFAF"

    # Le modèle livré avec python-docx omet un attribut obligatoire du zoom.
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    rendu = RenduWord(doc)
    construire(rendu)
    doc.save(OUT)
    print(f"{OUT} — {rendu.n_fig} figures, {rendu.n_tab} tableaux")


if __name__ == "__main__":
    main()
